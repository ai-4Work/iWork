"""生成复杂任务测试用例 Excel。

复杂 = 10–20 次工具调用组成（见 docs/chapters/15-测试评估逻辑.md 第2层）。本表取判据的「调用数」那一支，
把工具面从简单档的 1 个、中度档的 2–4 个抬到 6–8 个：一条链串起 6 个以上工具/skill 才算复杂。
判据的另一支（重特征：上下文压缩 / 子 agent / recall 召回 / 超时对账）本轮不铺——四项要真正立起来都得动测试装置
（团队会话 fixture / 把 330s 超时改小 / seed offloaded_blocks / 刮服务端日志），不属于「写用例」这一层；
只在 C9 放了一个 write_memory + recall 的轻量占位。

列结构与样式复刻 docs/chapters/中度任务-内置工具+Skill测试用例.xlsx，仅三处按内容放宽：
被测组合列宽 22→34、参考轨迹列宽 60→92、数据行高 60→110（要装 10–20 步）。

计数口径：计入 10–20 额度的是**客户端工具往返**（bash / read_file / write_file / edit_file / glob / grep）。
skill(name=…) 只把 SKILL.md 文本注入上下文、不产生客户端往返，故在「参考轨迹」里标出但不计数；
skill 的功能通过 bash 跑其脚本，计 1 次。服务端工具（write_memory / recall）同样不计数，但计入「被测组合」的面数。

用法:
    python i-work/tools/gen_complex_testcases.py

各条命令串均已对 i-work/server/skills/definitions/ 下的真实脚本核对（脚本名、参数、输出位置）；
凡真实行为与 SKILL.md 不符之处，已写进「期望行为」第二段。本档新核 / 复用的要点:
  - paddleocr：只有 parse_pdf_direct.py 能正确解析现网 API（内容嵌在
    result.result.layoutParsingResults[].markdown.text）；parse_document.py / extract_fields.py
    读顶层 result["layout"] 等键，输出近乎为空、字段全 null。--output 在 parse_pdf_direct.py 上是**目录**、
    页文件按 page_N.md 命名（另外两个脚本上是文件路径）。
  - editorial-diagrams：只有 scripts/validate_svg.py，无任何 flag（--help 会被当文件名）；
    ALLOWED_FONT_SIZES = {8,12,16,20,24,28,32,40} 把 SKILL.md §4.3 要求的 sublabel 9px 判为 ERROR ——
    照文档写必然校验失败。遮罩问题只 WARN、不阻断。
  - obsidian-cli：无 scripts/ 目录，全靠 PATH 上的 obsidian 二进制（1.12+，须已运行且开了 CLI，headless 会失败）；
    参数是 key=value、vault= 必须在子命令之前（是所有参数里的第一个）；read file=<无扩展名 wikilink>
    与 read path="…​.md" 两种写法不可混用。
  - image-processor：格式转换的 flag 是 --convert（不是 --format）；--compress 在 Pillow 路径下只对 JPEG 输出生效、
    PNG 上等价空操作；默认输出名 <stem>_processed.<ext>；手写 argv 解析（不是 argparse）。
  - wecom-unified：主力是外部 wecom-cli（≥1.1.0，须已完成扫码授权；auth init --noninteractive 会阻塞等扫码，
    测试里不可用）；build_docx.py 无 --output、输出固定落 {WRITABLE_ROOT}/docx/{spec_stem}.docx，
    其路径只在 references 里、SKILL.md 正文没提。「输出不得含原始 ID」是纯提示词约束，脚本层无强制。
  - imap-smtp-email：手写 parseArgs()，未知 flag 静默接受、不报错；--account 被 config.js 先从 argv 剥掉（全局参数）；
    download 落盘用附件自带文件名（不可预知）；结果 JSON 到 stdout、错误到 stderr。
  - smart-charts：产出的 HTML 文件名是 {safe_title[:30]}_{md5前6}.html，后缀随内容变、不可预知；
    把 assets/echarts.min.js 内联进 HTML，离线自包含、无 CDN 外链；cli.py 是 argparse，形如
    `python scripts/cli.py <file> <chart_type> --output-dir out`。
  - browser-skill：只有 SKILL.md、无 scripts；命令是 bsk session start / navigate <url> --session <id> /
    observe --session <id> / get-html --session <id> / screenshot --out <path> / session stop <id>；
    **没有 get-text**（该命令不存在）；start / stop 必须成对，stop 在成功与失败路径都要跑。
  - docx 用 docx-js（`npm install -g docx`）、pptx 用 pptxgenjs（`npm install -g pptxgenjs`）从零生成，
    产物分别过 scripts/office/validate.py 与 scripts/office/unpack.py。
  - summarize 的 --length 默认即 medium；产物默认只到 stdout，需要落盘时由 agent 自己 write_file。
  - pdf 的表格抽取按 SKILL.md 推荐走 pdfplumber 的 page.extract_tables()。
  - weekly-report 是纯提示词技能（零代码），四段结构 + 明令禁止凭空编造。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "用例编号",
    "被测组合",
    "用例名称",
    "前置条件（Fixture）",
    "输入 Prompt",
    "期望行为",
    "参考轨迹（完整序列）",
    "协作特征",
    "运行层",
]

COLUMN_WIDTHS = [9, 34, 16, 30, 28, 58, 92, 20, 14]

# 居中列：A 用例编号、I 运行层；其余左对齐
CENTER_COLUMNS = {1, 9}

OUTPUT_NAME = "复杂任务-内置工具+Skill测试用例.xlsx"

SHEET_NAME = "复杂任务"


CASE_DATA: list[list[str]] = [
    # ── C1 ────────────────────────────────────────────────────────────────
    [
        "C1",
        "glob + grep + read_file + summarize + xlsx + smart-charts + pptx",
        "竞品调研素材 → 汇总表 → 图 → 交付 PPT",
        "research/ 下 5 篇 .md 调研稿（各 3–15KB，含「竞品」小节）；无 out/；已装 pptxgenjs",
        "把 research/ 里跟竞品相关的素材汇总成一张对比表，出张柱状图，最后做一份 PPT",
        "产出 competitors.xlsx、一张柱状图 HTML 与 out/deck.pptx 三件套，PPT 引用的数据与表、图一致\n"
        "三个 read_file 同处一个读段（并发下发、未串行试探）；汇总只取「竞品」小节、不是把 5 篇全文塞进上下文；"
        "PPT 用 pptxgenjs 从零生成、产物能被 scripts/office/unpack.py 正常解开（不是改写模板）；"
        "out/ 下的 HTML 内联 echarts.min.js、不含 src=\"http 外链；summarize 走 summarize.py 且 --length long；"
        "write_file(build_deck.js) 单独成写段、不与任何只读调用并发",
        "skill(summarize) → glob(research/*.md) → grep(竞品, research/) "
        "→ [读段] read_file(research/a.md) ∥ read_file(research/b.md) ∥ read_file(research/c.md) "
        "→ bash(python summarize.py --file research/a.md --length long --output markdown) "
        "→ skill(xlsx) → bash(python pandas: 汇总三篇要点 → competitors.xlsx) "
        "→ skill(smart-charts) → bash(python scripts/cli.py competitors.xlsx bar --output-dir out) "
        "→ [写段] write_file(build_deck.js) → skill(pptx) "
        "→ bash(node build_deck.js → out/deck.pptx) "
        "→ bash(python scripts/office/unpack.py out/deck.pptx out/unpacked) → 结论",
        "读并发 / 大文件不整读 / 脚本执行 / 写串行",
        "slow + skill",
    ],
    # ── C2 ────────────────────────────────────────────────────────────────
    [
        "C2",
        "imap-smtp-email + pdf + xlsx + smart-charts + docx + write_file",
        "邮件取附件 → 抽表 → 双图 → 月度 Word 报告",
        "已配 ~/.config/mail-skills/.env 且联网；收件箱有两封带 PDF 附件的对账邮件；已装 pdfplumber、npm docx、soffice",
        "把最近两封对账邮件的附件表合到一起，出两张图（柱状 + 折线），再写份月度报告 Word",
        "产出 monthly.xlsx、两张图 HTML 与 monthly.docx\n"
        "附件落盘名取自邮件本身、不可预知——先 list / glob 再取，不按固定名找；"
        "两张图是两次 scripts/cli.py 调用（bar / line 各一），不是一张图凑两处引用；"
        "monthly.docx 通过 scripts/office/validate.py；报告里的数字与 xlsx、图的 series 一致（不得编造）；"
        "写段（write_file(build_report.js)）独立成段、不与只读调用并发下发",
        "skill(imap-smtp-email) → bash(node scripts/imap.js search --since <date> --limit 5) "
        "→ bash(node scripts/imap.js download <uid1> --dir attachments) "
        "→ bash(node scripts/imap.js download <uid2> --dir attachments) "
        "→ skill(pdf) → bash(python pdfplumber: 抽表 → out/monthly.csv) "
        "→ skill(xlsx) → bash(python pandas: → monthly.xlsx) "
        "→ skill(smart-charts) → bash(python scripts/cli.py monthly.xlsx bar --output-dir out) "
        "→ bash(python scripts/cli.py monthly.xlsx line --output-dir out) "
        "→ [写段] write_file(build_report.js) → skill(docx) "
        "→ bash(node build_report.js → monthly.docx) "
        "→ bash(python scripts/office/validate.py monthly.docx) "
        "→ read_file(out 下柱状图 HTML 复核) → 结论",
        "外部 CLI / 批量幂等 / 脚本执行 / 写串行",
        "slow + skill",
    ],
    # ── C3 ────────────────────────────────────────────────────────────────
    [
        "C3",
        "glob + grep + paddleocr + summarize + xlsx + obsidian-cli + write_file",
        "扫描件 OCR → 索引表 → 摘要 → 归档笔记库",
        "scans/ 下 4 份图片型扫描 PDF（正文不可直接抽文本）；已配 PADDLEOCR_API_URL / ACCESS_TOKEN 且联网；Obsidian 1.12+ 已运行且开了 CLI",
        "把 scans 里的扫描件识别出来做成一张索引表，挑一份摘要一下，然后归档到我的笔记库",
        "out/ocr/ 下逐页 markdown + index.xlsx + vault 里一条 ScanIndex 笔记 + audit.md\n"
        "只调 scripts/parse_pdf_direct.py，不碰 parse_document.py / extract_fields.py"
        "（后两者对现网 API 输出近乎为空、字段全 null）；该命令的 --output 是目录语义、页文件按 page_N.md 命名；"
        "原文 PDF 不整份读进上下文（只读 OCR 产物）；obsidian 的 vault= 必须紧跟 obsidian、在所有参数之前；"
        "index.xlsx 行数 = 实际页数，不虚增；write_file(audit.md) 单独成写段",
        "glob(scans/*.pdf) → skill(paddleocr) "
        "→ bash(for f in scans/*.pdf; do python scripts/parse_pdf_direct.py \"$f\" --output out/ocr --format markdown; done) "
        "→ bash(ls out/ocr/) → read_file(out/ocr/page_1.md) → grep(金额, out/ocr/) "
        "→ skill(summarize) → bash(python summarize.py --file out/ocr/page_1.md --length medium --output markdown) "
        "→ skill(xlsx) → bash(python pandas: 逐页登记 → index.xlsx) "
        "→ skill(obsidian-cli) → bash(obsidian vault=MyVault search query=\"2026\" format=json) "
        "→ bash(obsidian vault=MyVault create name=ScanIndex content=\"…\") "
        "→ [写段] write_file(audit.md) → bash(grep -c \"page_\" out/ocr/ 复核) → 结论",
        "批量幂等 / 大文件不整读 / 外部 CLI / 写串行",
        "slow + skill",
    ],
    # ── C4 ────────────────────────────────────────────────────────────────
    [
        "C4",
        "glob + grep + read_file + write_file + editorial-diagrams + obsidian-cli",
        "代码仓库 → 架构图 + 时序图 → 双笔记",
        "src/ 下 12 个 .py 模块；Obsidian 1.12+ 已运行且开了 CLI；vault 名已知",
        "看下 src 里的模块结构，画一张架构图和一张时序图，分别存到笔记库里",
        "out/arch.html 与 out/sequence.svg 各自通过 validate_svg.py（退出码 0），vault 里多出 Arch / Seq 两条笔记\n"
        "两张图是两个独立的写段，各自画出、各自校验，不合并成一次；"
        "字号只取 {8,12,16,20,24,28,32,40}——SKILL.md §4.3 要求 sublabel 用 9px，"
        "但校验器 ALLOWED_FONT_SIZES 把 9 判为 ERROR，照文档写必然失败；"
        "两处 create 的 content 引用图名、vault= 均在子命令之前；三个 read_file 并发、先用 grep 筛出定义再读",
        "glob(src/**/*.py) → grep(\"^def |^class \", src/) "
        "→ [读段] read_file(src/a.py) ∥ read_file(src/b.py) ∥ read_file(src/c.py) "
        "→ [写段] write_file(out/arch.html) → skill(editorial-diagrams) "
        "→ bash(python scripts/validate_svg.py out/arch.html) "
        "→ [写段] write_file(out/sequence.svg) "
        "→ bash(python scripts/validate_svg.py out/sequence.svg) "
        "→ skill(obsidian-cli) → bash(obsidian vault=MyVault create name=Arch content=\"![[arch.html]]\") "
        "→ bash(obsidian vault=MyVault create name=Seq content=\"![[sequence.svg]]\") → 结论",
        "读并发 / 外部 CLI / 写串行 / 幂等重放",
        "slow + skill",
    ],
    # ── C5 ────────────────────────────────────────────────────────────────
    [
        "C5",
        "glob + read_file + image-processor + pptx + docx + write_file",
        "图片资产三连处理 → PPT + Word 双交付",
        "assets/ 下 6 张图（4 jpg + 2 png）与 manifest.md；已装 Pillow、soffice、pptxgenjs、npm docx",
        "把 assets 里的图都压一下、加上水印、统一转成 jpg，然后做份 PPT 和一份说明 Word",
        "out/assets_deck.pptx 与 out/assets.docx 双交付，且分别通过 unpack.py / validate.py\n"
        "三次批处理各只跑一条循环命令（不是 N 张图 N 条命令）；格式转换的 flag 是 --convert、不是 --format；"
        "--compress 在 Pillow 路径下只对 JPEG 输出生效、PNG 上等价空操作，PNG 的压缩结果与输入一致属预期；"
        "默认输出名是 <stem>_processed.<ext>、process-image.py 是手写 argv 解析（不是 argparse）；"
        "两个交付物是两个独立写段，不共用一个 JS 文件",
        "glob(assets/*) → read_file(assets/manifest.md) → skill(image-processor) "
        "→ bash(for f in assets/*.jpg; do python scripts/process-image.py \"$f\" --compress 70 --output out/; done) "
        "→ bash(for f in assets/*.png; do python scripts/process-image.py \"$f\" --watermark \"ACME\" --output out/; done) "
        "→ bash(for f in out/*.png; do python scripts/process-image.py \"$f\" --convert jpg --output out/; done) "
        "→ [写段] write_file(build_deck.js) → skill(pptx) "
        "→ bash(node build_deck.js → out/assets_deck.pptx) "
        "→ bash(python scripts/office/unpack.py out/assets_deck.pptx out/unpacked) "
        "→ [写段] write_file(build_doc.js) → skill(docx) "
        "→ bash(node build_doc.js → out/assets.docx) "
        "→ bash(python scripts/office/validate.py out/assets.docx) → 结论",
        "批量幂等 / 脚本执行 / 写串行",
        "slow + skill",
    ],
    # ── C6 ────────────────────────────────────────────────────────────────
    [
        "C6",
        "glob + read_file + xlsx + smart-charts + docx + write_file",
        "多源表合并清洗 → 三张图 → 看板 Word",
        "data/ 下 3 个季度表 .xlsx（列名有出入、含重复行与空值）；已装 pandas / openpyxl、npm docx",
        "把 data 里三个季度的表合并清洗一下，出三张图，再做个数据看板 Word",
        "out/merged.xlsx、out/clean.xlsx、三张图 HTML 与 out/dashboard.docx\n"
        "三个 read_file 并发（不逐个串行试探）；合并与清洗分成两步、各自落盘，"
        "清洗前后的行数变化有据可查（去重与补缺各去/补了多少）；三张图是三次 scripts/cli.py 调用、三种图型；"
        "dashboard.docx 通过 scripts/office/validate.py，正文引用的数字与 clean.xlsx、图的 series 一致（不得编造）",
        "glob(data/*.xlsx) "
        "→ [读段] read_file(data/q1.xlsx) ∥ read_file(data/q2.xlsx) ∥ read_file(data/q3.xlsx) "
        "→ skill(xlsx) → bash(python pandas: concat → out/merged.xlsx) "
        "→ bash(python pandas: 去重补缺 → out/clean.xlsx) "
        "→ skill(smart-charts) → bash(python scripts/cli.py out/clean.xlsx line --output-dir out) "
        "→ bash(python scripts/cli.py out/clean.xlsx bar --output-dir out) "
        "→ bash(python scripts/cli.py out/clean.xlsx pie --output-dir out) "
        "→ [写段] write_file(build_dash.js) → skill(docx) "
        "→ bash(node build_dash.js → out/dashboard.docx) "
        "→ bash(python scripts/office/validate.py out/dashboard.docx) → 结论",
        "读并发 / 脚本执行 / 写串行",
        "slow + skill",
    ],
    # ── C7 ────────────────────────────────────────────────────────────────
    [
        "C7",
        "imap-smtp-email + pdf + obsidian-cli + weekly-report + docx + wecom-unified",
        "邮件 + 日记 → 周报 Word → 企微发出",
        "已配 IMAP 凭据与 .env；收件箱有一封带 PDF 附件的进展邮件；Obsidian 1.12+ 已运行且开了 CLI、有本周日记；wecom-cli 已完成扫码授权",
        "结合这周的邮件和我的日记，写份周报发到企业微信给我",
        "weekly.docx 通过 validate.py，并成功经 wecom-cli 发出\n"
        "周报只写邮件与日记里确实有的内容——weekly-report 明令禁止凭空编造，四段结构齐备；"
        "发出命令的输出不得回显 chat_id 等原始 ID；obsidian 的 vault= 在子命令之前、read file= 用无扩展名 wikilink 写法"
        "（不与 read path=\"…​.md\" 混用）；write_file(weekly.md) 单独成写段",
        "skill(imap-smtp-email) → bash(node scripts/imap.js search --since <date> --limit 20) "
        "→ bash(node scripts/imap.js download <uid> --dir attachments) "
        "→ skill(pdf) → bash(python pdfplumber: 抽要点 → out/attach.txt) "
        "→ skill(obsidian-cli) → bash(obsidian vault=MyVault search query=\"日记\" format=json limit=7) "
        "→ bash(obsidian vault=MyVault read file=日记) "
        "→ read_file(out/attach.txt) → skill(weekly-report) → [提示词生成四段式周报] "
        "→ [写段] write_file(weekly.md) → skill(docx) → bash(node docx-js: sections → weekly.docx) "
        "→ bash(python scripts/office/validate.py weekly.docx) "
        "→ skill(wecom-unified) → bash(wecom-cli message aibot send --json '{…}') → 结论",
        "外部 CLI / 提示词 / 有始有终 / 写串行",
        "slow + skill",
    ],
    # ── C8 ────────────────────────────────────────────────────────────────
    [
        "C8",
        "browser-skill + image-processor + xlsx + smart-charts + editorial-diagrams + write_file",
        "网页抓表 → 出图 + 压缩 → 流程图 → 落盘",
        "已装 bsk CLI 与浏览器扩展、Chromium 已登录；目标页有一张 ≥10 行的表；已装 Pillow / pandas",
        "把这个网页上的表格抓下来出张图，顺手把页面截图压一下，再画个流程图标一下步骤",
        "out/page.html、page.xlsx、柱状图 HTML、out/flow.svg 四件产物\n"
        "bsk session start / stop 成对，成功与失败路径都要 stop；导航后先 observe 确认元素、不猜选择器"
        "（browser-skill 没有 get-text 这个命令，只有 observe / get-html）；"
        "页面 HTML 先落盘再由 pandas 解析，不把整页 HTML 塞进上下文；"
        "out/flow.svg 通过 validate_svg.py（字号只取 {8,12,16,20,24,28,32,40}）",
        "skill(browser-skill) → bash(bsk session start) → bash(bsk navigate <url> --session <id>) "
        "→ bash(bsk observe --session <id>) → bash(bsk get-html --session <id> > out/page.html) "
        "→ bash(bsk session stop <id>) → skill(image-processor) "
        "→ bash(for f in out/page*.png; do python scripts/process-image.py \"$f\" --compress 70 --output out/; done) "
        "→ skill(xlsx) → bash(python pandas: 解析 out/page.html 表 → page.xlsx) "
        "→ skill(smart-charts) → bash(python scripts/cli.py page.xlsx bar --output-dir out) "
        "→ [写段] write_file(out/flow.svg) → skill(editorial-diagrams) "
        "→ bash(python scripts/validate_svg.py out/flow.svg) → 结论",
        "外部 CLI / 有始有终 / 大文件不整读 / 写串行",
        "slow + skill",
    ],
    # ── C9 ────────────────────────────────────────────────────────────────
    [
        "C9",
        "glob + grep + bash + read_file + edit_file + write_file + write_memory + recall",
        "日志巡检 → 定位配置 → 改 → 出报告 → 记进记忆",
        "logs/ 下 3 个 .log（app.log 40MB、error.log、access.log）；conf/ 下 app.ini 与 limits.ini",
        "查一下最近的错误日志，看是不是连接池配置的问题，该改就改，再出份巡检报告",
        "conf/app.ini 的 pool_size 由 5 改为 20、audit_report.md 落盘、一条巡检结论写入记忆\n"
        "全文不出现 read_file(任意 .log)——40MB 的 app.log 靠 grep -c / tail 抽样，不整读；"
        "edit_file 单独成写段、old_string 在文件里唯一，改完用 grep -n 复核；"
        "write_memory / recall 是服务端工具，走 tool_use 但不产生客户端往返（不计入 10–20 额度，只计面数）；"
        "报告里的错误计数与第 2 步的 grep -c 输出一致",
        "glob(logs/*.log) "
        "→ bash(for f in logs/*.log; do printf '%s ' \"$f\"; grep -c ERROR \"$f\"; done) "
        "→ bash(tail -n 100 logs/app.log) "
        "→ [读段] read_file(conf/app.ini) ∥ read_file(conf/limits.ini) "
        "→ bash(python: 按 limits.ini 阈值比对 app.ini 的 pool_size) "
        "→ [写段] edit_file(conf/app.ini: pool_size=5 → 20) "
        "→ bash(grep -n \"pool_size\" conf/app.ini 复核) "
        "→ write_memory(日志巡检结论) → recall(query=\"池大小约定\") "
        "→ [写段] write_file(audit_report.md) "
        "→ bash(grep -c \"ERROR\" logs/error.log 复核报告数字) → 结论",
        "大文件不整读 / 失败恢复 / 记忆写入 / 记忆召回 / 写串行",
        "slow + skill",
    ],
    # ── C10 ───────────────────────────────────────────────────────────────
    [
        "C10",
        "imap-smtp-email + glob + xlsx + smart-charts + editorial-diagrams + docx + wecom-unified",
        "邮件附件 vs ERP 对账 → 差异图 → 报告 → 企微",
        "已配 IMAP 凭据与 .env；收件箱有一封带对账 CSV 附件的邮件；workspace 有 erp_export.xlsx；wecom-cli 已完成扫码授权",
        "把邮件里的对账表和我们的 ERP 导出一比，差异画张图，出份报告发到企业微信",
        "out/diff.xlsx、差异图 HTML、out/diff_diagram.svg、diff_report.docx，并经 wecom-cli 发出\n"
        "附件落盘名取自邮件本身、不可预知（先 glob 再读）；差异表的行数 = 实际不一致条数，不凑数、不截断；"
        "diff.xlsx、out/diff_diagram.svg、diff_report.docx 三处产物分别过各自校验"
        "（后两者走 validate_svg.py / validate.py）；发出命令的输出不含原始 ID；两个写段各自独立",
        "skill(imap-smtp-email) → bash(node scripts/imap.js check --limit 20) "
        "→ bash(node scripts/imap.js download <uid> --dir attachments) "
        "→ glob(attachments/*) → read_file(attachments/<name>.csv) "
        "→ skill(xlsx) → bash(python pandas: ERP 表 vs 邮件表比对 → out/diff.xlsx) "
        "→ skill(smart-charts) → bash(python scripts/cli.py out/diff.xlsx bar --output-dir out) "
        "→ [写段] write_file(out/diff_diagram.svg) → skill(editorial-diagrams) "
        "→ bash(python scripts/validate_svg.py out/diff_diagram.svg) "
        "→ [写段] write_file(build_diff_doc.js) → skill(docx) "
        "→ bash(node build_diff_doc.js → diff_report.docx) "
        "→ bash(python scripts/office/validate.py diff_report.docx) "
        "→ skill(wecom-unified) → bash(wecom-cli message aibot send --json '{…}') → 结论",
        "外部 CLI / 大文件不整读 / 脚本执行 / 写串行",
        "slow + skill",
    ],
]


HEADER_FILL = PatternFill("solid", fgColor="FF1F4E79")
DATA_FILL = PatternFill("solid", fgColor="FFF2F7FB")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 本档唯一的结构性偏离：10–20 步的轨迹要更高的行（中度档是 60）
DATA_ROW_HEIGHT = 110


def apply_style(ws, num_data_rows: int) -> None:
    """复刻源表样式：表头深蓝白字、数据行浅蓝、细边框、列宽、行高、冻结表头。"""
    for idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name="宋体", size=11, bold=True, color="FFFFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 26

    for row in range(2, num_data_rows + 2):
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(
                name="Calibri",
                size=10,
                bold=(col == 1),
            )
            cell.fill = DATA_FILL
            cell.alignment = Alignment(
                horizontal="center" if col in CENTER_COLUMNS else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = BORDER
        ws.row_dimensions[row].height = DATA_ROW_HEIGHT

    ws.freeze_panes = "A2"


def build(output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(title=SHEET_NAME)
    ws.append(HEADERS)
    for row in CASE_DATA:
        ws.append(row)
    apply_style(ws, len(CASE_DATA))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def count_client_calls(trajectory: str) -> int:
    """数一遍轨迹里的客户端工具往返（skill(...) / write_memory / recall 不计）。"""
    import re

    return len(re.findall(r"\b(?:read_file|write_file|edit_file|glob|grep|bash)\(", trajectory))


def main() -> None:
    base = Path(__file__).resolve().parent.parent
    output_path = base / "docs" / "chapters" / OUTPUT_NAME
    build(output_path)

    print(f"written: {output_path}")
    for row in CASE_DATA:
        calls = count_client_calls(row[6])
        faces = len(row[1].split(" + "))
        flag = "OK " if 10 <= calls <= 20 and faces >= 6 else "BAD"
        print(f"  [{flag}] {row[0]}  调用 {calls} 次  面 {faces} 个  {row[2]}")
    print(f"  total: {len(CASE_DATA)} 条")


if __name__ == "__main__":
    main()
